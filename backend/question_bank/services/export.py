"""
Printable exports for question papers: PDF (PyMuPDF Story, already a
dependency) and DOCX (python-docx). LaTeX is printed as written — teachers
who need typeset maths can edit the DOCX.
"""
from __future__ import annotations

import html
import io
import re

from ..models import PaperItem, QuestionPaper

_LETTERS = "ABCDEF"

_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.35; }
h1 { font-size: 16pt; text-align: center; margin: 0 0 2pt 0; }
.sub { text-align: center; font-size: 10pt; color: #333; margin: 0 0 4pt 0; }
.meta { font-size: 10pt; margin: 6pt 0; }
.instr { font-size: 10pt; border: 1px solid #999; padding: 4pt; margin: 6pt 0; }
h2 { font-size: 12pt; margin: 12pt 0 2pt 0; border-bottom: 1px solid #999; }
.sec-instr { font-size: 9.5pt; font-style: italic; margin: 0 0 4pt 0; }
.q { margin: 6pt 0 2pt 0; }
.marks { float: right; font-size: 9pt; color: #444; }
.opt { margin: 1pt 0 0 18pt; font-size: 10.5pt; }
.ans { margin: 2pt 0 0 18pt; font-size: 10pt; color: #14532d; }
.exp { margin: 1pt 0 0 18pt; font-size: 9.5pt; color: #444; }
"""


def _esc(text: str) -> str:
    return html.escape(text or "").replace("\n", "<br/>")


def _answer_text(item: PaperItem) -> str:
    if item.qtype in ("mcq", "multi") and item.correct_indices:
        letters = ", ".join(_LETTERS[i] for i in item.correct_indices if 0 <= i < len(_LETTERS))
        return f"({letters})"
    return item.answer or "—"


def _sections(paper: QuestionPaper):
    for section in paper.sections.prefetch_related("items").all():
        yield section, list(section.items.all())


def paper_html(paper: QuestionPaper, *, include_answers: bool) -> str:
    total = paper.total_marks()
    parts = [f"<h1>{_esc(paper.title)}</h1>"]
    sub = " · ".join(x for x in [paper.exam_name, paper.class_label, paper.bucket.name] if x)
    if sub:
        parts.append(f'<p class="sub">{_esc(sub)}</p>')
    parts.append(
        f'<p class="meta"><b>Time:</b> {paper.duration_minutes} minutes'
        f' &nbsp;&nbsp; <b>Maximum marks:</b> {total}</p>'
    )
    if include_answers:
        parts.append('<p class="meta"><b>ANSWER KEY &amp; MARKING SCHEME</b></p>')
    if paper.instructions:
        parts.append(f'<div class="instr"><b>General instructions</b><br/>{_esc(paper.instructions)}</div>')

    n = 0
    for section, items in _sections(paper):
        if not items:
            continue
        sec_marks = sum(i.marks for i in items)
        parts.append(f"<h2>{_esc(section.title)} <span class='marks'>{sec_marks} marks</span></h2>")
        if section.instructions:
            parts.append(f'<p class="sec-instr">{_esc(section.instructions)}</p>')
        for item in items:
            n += 1
            parts.append(
                f'<p class="q"><b>Q{n}.</b> {_esc(item.prompt)}'
                f' <span class="marks">[{item.marks}]</span></p>'
            )
            for idx, opt in enumerate(item.options or []):
                letter = _LETTERS[idx] if idx < len(_LETTERS) else str(idx + 1)
                parts.append(f'<p class="opt">({letter}) {_esc(opt)}</p>')
            if include_answers:
                parts.append(f'<p class="ans"><b>Answer:</b> {_esc(_answer_text(item))}</p>')
                if item.explanation:
                    parts.append(f'<p class="exp">{_esc(item.explanation)}</p>')
    return "".join(parts)


def paper_pdf(paper: QuestionPaper, *, include_answers: bool) -> bytes:
    import fitz

    story = fitz.Story(html=paper_html(paper, include_answers=include_answers), user_css=_CSS)
    buffer = io.BytesIO()
    writer = fitz.DocumentWriter(buffer)
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (42, 42, -42, -42)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    return buffer.getvalue()


def paper_docx(paper: QuestionPaper, *, include_answers: bool) -> bytes:
    from docx import Document as Docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Docx()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    head = doc.add_heading(paper.title, level=1)
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = " · ".join(x for x in [paper.exam_name, paper.class_label, paper.bucket.name] if x)
    if sub:
        p = doc.add_paragraph(sub)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        f"Time: {paper.duration_minutes} minutes    Maximum marks: {paper.total_marks()}"
    )
    if include_answers:
        doc.add_paragraph().add_run("ANSWER KEY & MARKING SCHEME").bold = True
    if paper.instructions:
        p = doc.add_paragraph()
        p.add_run("General instructions\n").bold = True
        p.add_run(paper.instructions)

    n = 0
    for section, items in _sections(paper):
        if not items:
            continue
        doc.add_heading(f"{section.title} ({sum(i.marks for i in items)} marks)", level=2)
        if section.instructions:
            doc.add_paragraph(section.instructions).runs[0].italic = True
        for item in items:
            n += 1
            p = doc.add_paragraph()
            p.add_run(f"Q{n}. ").bold = True
            p.add_run(item.prompt)
            p.add_run(f"  [{item.marks}]")
            for idx, opt in enumerate(item.options or []):
                letter = _LETTERS[idx] if idx < len(_LETTERS) else str(idx + 1)
                doc.add_paragraph(f"({letter}) {opt}").paragraph_format.left_indent = Pt(18)
            if include_answers:
                a = doc.add_paragraph()
                a.paragraph_format.left_indent = Pt(18)
                a.add_run("Answer: ").bold = True
                a.add_run(_answer_text(item))
                if item.explanation:
                    e = doc.add_paragraph(item.explanation)
                    e.paragraph_format.left_indent = Pt(18)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def export_filename(paper: QuestionPaper, ext: str, *, include_answers: bool) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "-", paper.title).strip("-")[:60] or "question-paper"
    suffix = "-answer-key" if include_answers else ""
    return f"{base}{suffix}.{ext}"
