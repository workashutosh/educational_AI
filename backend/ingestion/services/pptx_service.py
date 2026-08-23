"""
Generate a downloadable PPTX from bucket notes (RAG + chat LLM outline + python-pptx).

Supports multiple visual styles and creative slide layouts (hooks, comparisons, quizzes).
"""
from __future__ import annotations

import io
import json
import logging
import re
from dataclasses import dataclass

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .llm_provider import chat_completion, get_chat_llm
from .retrieval_service import search_bucket

logger = logging.getLogger(__name__)

VALID_STYLES = frozenset({"creative", "clean", "bold", "sunset", "ocean"})
VALID_LAYOUTS = frozenset({"standard", "hook", "compare", "quiz", "quote", "summary"})

_BASE_RULES = (
    "- Produce exactly the requested number of slides in the slides array.\n"
    "- First slide should hook the audience; last should summarize key takeaways.\n"
    "- 3–6 short bullets per slide; no long paragraphs.\n"
    "- Prefer plain text; for formulas use simple Unicode or ASCII math "
    "(e.g. ∫ f(x) dx), not LaTeX dollar delimiters.\n"
    "- Speaker notes are optional and brief.\n"
    '- Each slide must include "layout" (see schema).'
)

_SYSTEM_CLEAN = (
    "You create clear teaching PowerPoint outlines for Class XII students.\n"
    "Use ONLY the numbered excerpts from the teacher's indexed notes. "
    "Do not invent facts that are not supported by the excerpts.\n\n"
    "Respond with valid JSON only (no markdown fences). Schema:\n"
    "{\n"
    '  "title": "deck title",\n'
    '  "slides": [\n'
    "    {\n"
    '      "heading": "slide title",\n'
    '      "subtitle": "optional short tagline",\n'
    '      "layout": "standard",\n'
    '      "bullets": ["point", "..."],\n'
    '      "callout": "optional key insight",\n'
    '      "notes": "optional speaker notes"\n'
    "    }\n"
    "  ]\n"
    "}\n\n"
    "Rules:\n"
    + _BASE_RULES
)

_SYSTEM_CREATIVE = (
    "You create engaging, visually varied teaching PowerPoint outlines for Class XII.\n"
    "Use ONLY the numbered excerpts. Ground every fact in the notes — never invent content.\n"
    "Make each deck feel fresh: vary tone, structure, and slide types so no two decks look alike.\n\n"
    "Respond with valid JSON only (no markdown fences). Schema:\n"
    "{\n"
    '  "title": "memorable deck title",\n'
    '  "slides": [\n'
    "    {\n"
    '      "heading": "slide title",\n'
    '      "subtitle": "punchy one-line hook (optional)",\n'
    '      "layout": "standard" | "hook" | "compare" | "quiz" | "quote" | "summary",\n'
    '      "bullets": ["point", "..."],\n'
    '      "callout": "highlight box text — analogy, fun fact, or key formula",\n'
    '      "notes": "optional speaker notes"\n'
    "    }\n"
    "  ]\n"
    "}\n\n"
    "Layout guide:\n"
    "- hook: open with a surprising question or real-world scenario (use callout).\n"
    "- compare: contrast two ideas side-by-side (even number of bullets helps).\n"
    "- quiz: quick check-for-understanding (callout = the question).\n"
    "- quote: memorable one-liner or mnemonic in callout; bullets = context.\n"
    "- summary: recap slide with actionable takeaways.\n"
    "- standard: teach a concept clearly.\n\n"
    "Creative rules:\n"
    "- Use at least 3 different layouts across the deck.\n"
    "- Do NOT repeat the same layout on consecutive slides.\n"
    "- Add subtitles or callouts on most slides — make it feel designed, not generic.\n"
    "- Use vivid but accurate headings; avoid boring labels like 'Introduction' alone.\n"
    + _BASE_RULES
)


class PptxGenerationError(Exception):
    """User-facing generation failure (empty retrieval, bad model output, etc.)."""


@dataclass(frozen=True)
class PptTheme:
    key: str
    label: str
    primary: RGBColor
    accent: RGBColor
    bg: RGBColor
    bg_alt: RGBColor
    callout_bg: RGBColor
    text: RGBColor
    text_muted: RGBColor
    title_size: int
    heading_size: int
    dark_title: bool = False


THEMES: dict[str, PptTheme] = {
    "clean": PptTheme(
        key="clean",
        label="Clean",
        primary=RGBColor(0x1A, 0x3A, 0x32),
        accent=RGBColor(0x0F, 0x76, 0x6E),
        bg=RGBColor(0xF7, 0xF9, 0xFD),
        bg_alt=RGBColor(0xEE, 0xF5, 0xF1),
        callout_bg=RGBColor(0xE0, 0xF2, 0xF1),
        text=RGBColor(0x22, 0x28, 0x26),
        text_muted=RGBColor(0x5A, 0x6B, 0x66),
        title_size=36,
        heading_size=28,
    ),
    "creative": PptTheme(
        key="creative",
        label="Creative",
        primary=RGBColor(0x1E, 0x1B, 0x4B),
        accent=RGBColor(0xE1, 0x1D, 0x48),
        bg=RGBColor(0xFA, 0xF5, 0xFF),
        bg_alt=RGBColor(0xFF, 0xF1, 0xF2),
        callout_bg=RGBColor(0xFE, 0xE2, 0xE2),
        text=RGBColor(0x1E, 0x1B, 0x4B),
        text_muted=RGBColor(0x6B, 0x72, 0x80),
        title_size=38,
        heading_size=30,
    ),
    "bold": PptTheme(
        key="bold",
        label="Bold",
        primary=RGBColor(0xFF, 0xFF, 0xFF),
        accent=RGBColor(0xF5, 0x9E, 0x0B),
        bg=RGBColor(0x11, 0x18, 0x27),
        bg_alt=RGBColor(0x1F, 0x29, 0x37),
        callout_bg=RGBColor(0x37, 0x47, 0x5A),
        text=RGBColor(0xF9, 0xFA, 0xFB),
        text_muted=RGBColor(0xD1, 0xD5, 0xDB),
        title_size=40,
        heading_size=32,
        dark_title=True,
    ),
    "sunset": PptTheme(
        key="sunset",
        label="Sunset",
        primary=RGBColor(0x7C, 0x2D, 0x12),
        accent=RGBColor(0xEA, 0x58, 0x0C),
        bg=RGBColor(0xFF, 0xF7, 0xED),
        bg_alt=RGBColor(0xFF, 0xED, 0xD5),
        callout_bg=RGBColor(0xFE, 0xD7, 0xAA),
        text=RGBColor(0x43, 0x2A, 0x1A),
        text_muted=RGBColor(0x9A, 0x34, 0x12),
        title_size=36,
        heading_size=28,
    ),
    "ocean": PptTheme(
        key="ocean",
        label="Ocean",
        primary=RGBColor(0x0C, 0x4A, 0x6E),
        accent=RGBColor(0x0E, 0xA5, 0xE9),
        bg=RGBColor(0xF0, 0xF9, 0xFF),
        bg_alt=RGBColor(0xE0, 0xF2, 0xFE),
        callout_bg=RGBColor(0xBA, 0xE6, 0xFD),
        text=RGBColor(0x0C, 0x4A, 0x6E),
        text_muted=RGBColor(0x03, 0x71, 0x9A),
        title_size=36,
        heading_size=28,
    ),
}


@dataclass
class GeneratedOutline:
    title: str
    slides: list[dict]
    model: str
    prompt: str
    style: str


@dataclass
class GeneratedPptx:
    data: bytes
    filename: str
    title: str
    slide_count: int


def generate_pptx_outline(
    bucket_id: int,
    prompt: str,
    slide_count: int = 8,
    top_k: int = 12,
    *,
    style: str = "creative",
) -> GeneratedOutline:
    """RAG + chat LLM → structured slide outline for preview (no file yet)."""
    prompt = (prompt or "").strip()
    if not prompt:
        raise PptxGenerationError("Prompt is required.")
    if slide_count < 3 or slide_count > 15:
        raise PptxGenerationError("slide_count must be between 3 and 15.")
    style = _normalize_style(style)

    cfg = get_chat_llm()
    model = cfg.model
    chunks = search_bucket(bucket_id, query=prompt, top_k=top_k, kinds=["notes"])
    if not chunks:
        raise PptxGenerationError(
            "No matching notes found in this AI Bucket. Index notes first, "
            "or try a more specific topic from your materials."
        )

    context = "\n\n".join(
        f"[{i}] (pages {c.page_start}-{c.page_end}, "
        f"{' > '.join(c.heading_path) or 'untitled section'}):\n{c.text}"
        for i, c in enumerate(chunks, start=1)
    )
    user_msg = (
        f"Excerpts:\n\n{context}\n\n"
        f"User request: {prompt}\n"
        f"Create exactly {slide_count} slides."
    )

    outline = _fetch_outline(user_msg, slide_count, style=style)
    return GeneratedOutline(
        title=outline["title"],
        slides=outline["slides"],
        model=model,
        prompt=prompt,
        style=style,
    )


def pptx_from_outline(
    title: str,
    slides: list[dict],
    *,
    style: str = "creative",
) -> GeneratedPptx:
    """Build a .pptx from a previewed outline (no LLM call)."""
    style = _normalize_style(style)
    outline = _normalize_outline({"title": title, "slides": slides}, slide_count=None)
    if not outline["slides"]:
        raise PptxGenerationError("Outline has no slides to export.")
    data = _build_pptx(outline, style=style)
    return GeneratedPptx(
        data=data,
        filename=_safe_filename(outline["title"]),
        title=outline["title"],
        slide_count=len(outline["slides"]),
    )


def generate_pptx(
    bucket_id: int,
    prompt: str,
    slide_count: int = 8,
    top_k: int = 12,
    *,
    style: str = "creative",
) -> GeneratedPptx:
    """Convenience: outline + build in one step."""
    outline = generate_pptx_outline(
        bucket_id, prompt=prompt, slide_count=slide_count, top_k=top_k, style=style
    )
    return pptx_from_outline(outline.title, outline.slides, style=style)


def _normalize_style(style: str | None) -> str:
    key = (style or "creative").strip().lower()
    return key if key in VALID_STYLES else "creative"


def _fetch_outline(user_msg: str, slide_count: int, *, style: str) -> dict:
    creative = style in ("creative", "sunset", "ocean")
    system = _SYSTEM_CREATIVE if creative else _SYSTEM_CLEAN
    temperature = 0.72 if creative else 0.35
    messages = [
        {"role": "system", "content": system},
        {"role": "user", "content": user_msg},
    ]
    raw = chat_completion(
        messages,
        temperature=temperature,
        max_tokens=4096,
        response_format={"type": "json_object"},
        timeout=90,
        retries_on_429=1,
    )
    data = _parse_outline_json(raw)
    return _normalize_outline(data, slide_count)


def _parse_outline_json(raw: str) -> dict:
    text = raw.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise PptxGenerationError(f"Model returned invalid JSON: {exc}") from exc
    if not isinstance(data, dict):
        raise PptxGenerationError("Model outline must be a JSON object.")
    return data


def _normalize_outline(data: dict, slide_count: int | None) -> dict:
    title = str(data.get("title") or "Lesson slides").strip() or "Lesson slides"
    slides_raw = data.get("slides")
    if not isinstance(slides_raw, list) or not slides_raw:
        raise PptxGenerationError("Model returned no slides.")

    slides: list[dict] = []
    for item in slides_raw:
        if not isinstance(item, dict):
            continue
        heading = str(item.get("heading") or item.get("title") or "Slide").strip()
        subtitle = str(item.get("subtitle") or "").strip()[:160]
        layout = str(item.get("layout") or "standard").strip().lower()
        if layout not in VALID_LAYOUTS:
            layout = "standard"
        bullets_in = item.get("bullets") or item.get("points") or []
        if isinstance(bullets_in, str):
            bullets_in = [bullets_in]
        bullets = [str(b).strip() for b in bullets_in if str(b).strip()][:8]
        if not bullets:
            body = str(item.get("content") or item.get("body") or "").strip()
            if body:
                bullets = [body]
        callout = str(item.get("callout") or "").strip()[:280]
        notes = str(item.get("notes") or "").strip()
        if not bullets and not callout:
            continue
        if not bullets:
            bullets = ["See speaker notes."]
        slides.append(
            {
                "heading": heading[:120],
                "subtitle": subtitle,
                "layout": layout,
                "bullets": bullets,
                "callout": callout,
                "notes": notes[:500],
            }
        )

    if not slides:
        raise PptxGenerationError("Model returned empty slide content.")

    if slide_count is not None:
        if len(slides) > slide_count:
            slides = slides[:slide_count]
        while len(slides) < slide_count:
            base = slides[-1]
            slides.append(
                {
                    "heading": f"{base['heading']} (cont.)",
                    "subtitle": "",
                    "layout": "standard",
                    "bullets": base["bullets"][:3],
                    "callout": "",
                    "notes": "",
                }
            )

    return {"title": title[:160], "slides": slides}


def _safe_filename(title: str) -> str:
    stem = re.sub(r"[^\w\s\-]+", "", title, flags=re.UNICODE).strip()
    stem = re.sub(r"\s+", "_", stem)[:60] or "lesson"
    return f"{stem}.pptx"


def _build_pptx(outline: dict, *, style: str) -> bytes:
    theme = THEMES[_normalize_style(style)]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    title_slide = prs.slides.add_slide(blank)
    _render_title_slide(title_slide, outline["title"], theme)

    for i, slide_data in enumerate(outline["slides"]):
        slide = prs.slides.add_slide(blank)
        _render_content_slide(slide, slide_data, theme, slide_index=i)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _render_title_slide(slide, title: str, theme: PptTheme) -> None:
    _fill_background(slide, theme.bg)
    _add_rect(slide, 0, 0, Inches(13.333), Inches(0.35), theme.accent)
    _add_rect(slide, Inches(0.55), Inches(1.8), Inches(0.12), Inches(3.2), theme.accent)
    _add_oval(slide, Inches(10.8), Inches(5.2), Inches(1.8), Inches(1.8), theme.bg_alt)

    title_color = theme.primary if not theme.dark_title else theme.primary
    _add_textbox(
        slide,
        Inches(0.95),
        Inches(2.0),
        Inches(10.5),
        Inches(1.8),
        title,
        font_size=theme.title_size,
        bold=True,
        color=title_color,
    )


def _render_content_slide(slide, data: dict, theme: PptTheme, *, slide_index: int) -> None:
    layout = data.get("layout") or "standard"
    alt = slide_index % 2 == 1
    _fill_background(slide, theme.bg_alt if alt else theme.bg)
    _add_rect(slide, 0, 0, Inches(13.333), Inches(0.22), theme.accent)
    _add_rect(slide, 0, Inches(0.22), Inches(0.14), Inches(7.28), theme.accent)

    top = Inches(0.55)
    if data.get("subtitle"):
        _add_textbox(
            slide,
            Inches(0.75),
            top,
            Inches(11.8),
            Inches(0.45),
            data["subtitle"],
            font_size=13,
            bold=True,
            color=theme.accent,
        )
        top = Inches(0.95)

    _add_textbox(
        slide,
        Inches(0.75),
        top,
        Inches(11.8),
        Inches(0.85),
        data["heading"],
        font_size=theme.heading_size,
        bold=True,
        color=theme.primary if not theme.dark_title else theme.text,
    )

    body_top = top + Inches(1.05)
    callout = (data.get("callout") or "").strip()

    if layout == "quote" and callout:
        _add_rounded_callout(
            slide,
            Inches(1.2),
            body_top,
            Inches(10.9),
            Inches(1.4),
            f'"{callout}"',
            theme,
            font_size=22,
            italic=True,
        )
        body_top += Inches(1.65)
    elif layout == "hook" and callout:
        _add_rounded_callout(
            slide,
            Inches(0.75),
            body_top,
            Inches(11.8),
            Inches(1.1),
            callout,
            theme,
            font_size=17,
            bold=True,
        )
        body_top += Inches(1.35)
    elif layout == "quiz" and callout:
        _add_rounded_callout(
            slide,
            Inches(0.75),
            body_top,
            Inches(11.8),
            Inches(0.95),
            f"Quick check: {callout}",
            theme,
            font_size=16,
            bold=True,
        )
        body_top += Inches(1.15)
    elif callout and layout == "standard":
        _add_rounded_callout(
            slide,
            Inches(0.75),
            body_top,
            Inches(11.8),
            Inches(0.85),
            callout,
            theme,
            font_size=14,
        )
        body_top += Inches(1.05)

    bullets = data.get("bullets") or []
    if layout == "compare" and len(bullets) >= 2:
        mid = (len(bullets) + 1) // 2
        left = bullets[:mid]
        right = bullets[mid:]
        _add_bullet_block(
            slide, Inches(0.85), body_top, Inches(5.5), Inches(4.8), left, theme
        )
        _add_bullet_block(
            slide, Inches(6.85), body_top, Inches(5.5), Inches(4.8), right, theme
        )
        _add_rect(slide, Inches(6.55), body_top, Inches(0.04), Inches(4.5), theme.accent)
    elif layout == "summary":
        lines = [f"✓ {b}" for b in bullets]
        _add_bullet_block(
            slide, Inches(0.85), body_top, Inches(11.5), Inches(5.0), lines, theme
        )
    else:
        prefix = "• "
        if layout == "quiz":
            prefix = "○ "
        lines = [f"{prefix}{b}" for b in bullets]
        _add_bullet_block(
            slide, Inches(0.85), body_top, Inches(11.5), Inches(5.0), lines, theme
        )

    if data.get("notes"):
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = data["notes"]


def _fill_background(slide, color: RGBColor) -> None:
    _add_rect(slide, 0, 0, Inches(13.333), Inches(7.5), color)


def _add_rect(slide, left, top, width, height, fill: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()


def _add_oval(slide, left, top, width, height, fill: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()


def _add_rounded_callout(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    theme: PptTheme,
    *,
    font_size: int = 15,
    bold: bool = False,
    italic: bool = False,
) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = theme.callout_bg
    shape.line.color.rgb = theme.accent
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER if italic else PP_ALIGN.LEFT
    font = p.runs[0].font
    font.size = Pt(font_size)
    font.bold = bold
    font.italic = italic
    font.name = "Calibri"
    font.color.rgb = theme.text


def _add_bullet_block(slide, left, top, width, height, lines: list[str], theme: PptTheme) -> None:
    text = "\n".join(lines)
    _add_textbox(
        slide,
        left,
        top,
        width,
        height,
        text,
        font_size=17,
        color=theme.text,
    )


def _add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    font_size: int = 18,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = None,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split("\n") or [""]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        if not p.runs:
            continue
        font = p.runs[0].font
        font.size = Pt(font_size)
        font.bold = bold
        font.italic = italic
        font.name = "Calibri"
        if color is not None:
            font.color.rgb = color
